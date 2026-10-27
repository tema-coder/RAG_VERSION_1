"""Тесты архитектурных сущностей BEAR и мапперов-адаптеров.

Проверяют новую архитектурную надстройку (Шаги 1–6 agent_promt.docx):
сущности ``entities.py``, маппер ``ParseResult -> ParsedDocument``,
маппер внутренних чанков в архитектурные, метод ``chunk_architectural`` и
новое поле API ``architectural_chunks``.

Старые тесты НЕ затрагивают — проверяют только новые сущности. Все
проверки из Шага 7 промпта:
  * ``chunk.id`` — валидный UUID;
  * ``chunk.metadata`` содержит ключи ``parent_id``, ``block_type``,
    ``breadcrumb``;
  * ``chunk.text`` не пустой.
"""

from __future__ import annotations

from uuid import UUID, uuid4

import pytest

from rag_parser.core.entities import (
    ContentBlock,
    ContentBlockType,
    ParsedDocument,
    Chunk as ArchitecturalChunk,
)
from rag_parser.core.interfaces import Chunk as InternalChunk, DocumentNode, ParseResult
from rag_parser.chunkers.chunk_mapper import (
    is_table,
    map_internal_to_architectural_chunk,
)
from rag_parser.parsers.result_mapper import map_parse_result_to_parsed_document


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #


def _internal_chunk(
    text: str = "Текст чанка",
    *,
    text_for_embedding: str | None = None,
    parent_id: str | None = "h_1",
    breadcrumb: str | None = "Заголовок",
    chunk_type: str | None = "text",
) -> InternalChunk:
    """Внутренний чанк с полями, как ставит HybridRagChunker."""
    metadata = {}
    if chunk_type is not None:
        metadata["chunk_type"] = chunk_type
    return InternalChunk(
        index=0,
        text=text,
        tokens=len(text.split()),
        source="test.docx",
        metadata=metadata,
        text_for_embedding=text_for_embedding,
        breadcrumb=breadcrumb,
        parent_id=parent_id,
        chunk_order=0,
    )


def _parse_result_with_tree() -> ParseResult:
    """ParseResult с деревом формы: заголовок -> параграф + таблица, + параграф без заголовка."""
    tree = DocumentNode(
        type="document",
        level=0,
        text="",
        children=[
            DocumentNode(
                type="header",
                level=1,
                text="Раздел 1. Введение",
                children=[
                    DocumentNode(
                        type="paragraph",
                        level=2,
                        text="Текст раздела 1.",
                        children=[],
                        metadata={},
                    ),
                    DocumentNode(
                        type="table",
                        level=2,
                        text="| A | B |\n|---|---|\n| 1 | 2 |",
                        children=[],
                        metadata={},
                    ),
                ],
                metadata={},
            ),
            DocumentNode(
                type="paragraph",
                level=0,
                text="Параграф без заголовка.",
                children=[],
                metadata={"annotations_count": 0},
            ),
        ],
        metadata={},
    )
    return ParseResult(
        document_id=uuid4().hex,
        filename="test.docx",
        format="dedoc",
        markdown="# Раздел 1. Введение\n\nТекст раздела 1.",
        tree=tree,
        metadata={"format": "dedoc"},
    )


# --------------------------------------------------------------------------- #
# Шаг 1: entities.py
# --------------------------------------------------------------------------- #


class TestEntities:
    """Архитектурные сущности BEAR (dataclass, UUID, enum)."""

    def test_content_block_type_values(self):
        assert ContentBlockType.TEXT.value == "text"
        assert ContentBlockType.TABLE.value == "table"
        assert ContentBlockType.IMAGE.value == "image"

    def test_content_block_construction(self):
        block = ContentBlock(
            type=ContentBlockType.TEXT,
            content="текст",
            position=0,
            metadata={"breadcrumb": "h"},
        )
        assert block.type is ContentBlockType.TEXT
        assert block.content == "текст"
        assert block.position == 0
        assert block.metadata["breadcrumb"] == "h"

    def test_parsed_document_construction(self):
        doc_id = uuid4()
        pd = ParsedDocument(
            document_id=doc_id,
            text="markdown",
            blocks=[],
            attachments=[],
            metadata={"k": "v"},
        )
        assert isinstance(pd.document_id, UUID)
        assert pd.document_id == doc_id
        assert pd.text == "markdown"
        assert pd.blocks == []
        assert pd.attachments == []
        assert pd.metadata == {"k": "v"}

    def test_architectural_chunk_construction(self):
        cid = uuid4()
        did = uuid4()
        chunk = ArchitecturalChunk(
            id=cid,
            document_id=did,
            document_version=1,
            sequence=0,
            text="текст эмбеддинга",
            token_count=3,
            metadata={"parent_id": "h_1"},
        )
        assert isinstance(chunk.id, UUID)
        assert isinstance(chunk.document_id, UUID)
        assert chunk.document_version == 1
        assert chunk.sequence == 0
        assert chunk.text == "текст эмбеддинга"
        assert chunk.metadata["parent_id"] == "h_1"


# --------------------------------------------------------------------------- #
# Шаг 2: result_mapper.py
# --------------------------------------------------------------------------- #


class TestResultMapper:
    """ParseResult -> ParsedDocument: обход дерева, типы блоков, breadcrumb."""

    def test_document_id_is_uuid(self):
        """parse_result.document_id (hex-str) -> ParsedDocument.document_id (UUID)."""
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        assert isinstance(pd.document_id, UUID)

    def test_text_is_markdown(self):
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        assert pd.text == "# Раздел 1. Введение\n\nТекст раздела 1."

    def test_attachments_is_empty_list(self):
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        assert pd.attachments == []

    def test_metadata_copied(self):
        pr = _parse_result_with_tree()
        pd = map_parse_result_to_parsed_document(pr)
        assert pd.metadata == {"format": "dedoc"}
        # копия, не ссылка
        pd.metadata["new"] = 1
        assert "new" not in pr.metadata

    def test_headers_do_not_become_blocks(self):
        """Заголовки НЕ создают отдельные ContentBlock (нужны для breadcrumb)."""
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        header_blocks = [
            b for b in pd.blocks if b.metadata.get("node_type") == "header"
        ]
        assert header_blocks == [], "заголовок не должен становиться блоком"

    def test_block_count_and_types(self):
        """Дерево: header->[paragraph, table], paragraph -> 3 блока (без header)."""
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        assert len(pd.blocks) == 3
        types = [b.type for b in pd.blocks]
        assert types == [
            ContentBlockType.TEXT,
            ContentBlockType.TABLE,
            ContentBlockType.TEXT,
        ]

    def test_breadcrumb_in_block_metadata(self):
        """Блоки под заголовком несут его текст в breadcrumb."""
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        # первые два блока — под «Раздел 1. Введение»
        assert pd.blocks[0].metadata["breadcrumb"] == "Раздел 1. Введение"
        assert pd.blocks[1].metadata["breadcrumb"] == "Раздел 1. Введение"
        # третий блок — без заголовка
        assert pd.blocks[2].metadata["breadcrumb"] is None

    def test_positions_are_sequential(self):
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        assert [b.position for b in pd.blocks] == [0, 1, 2]

    def test_table_content_is_markdown(self):
        pd = map_parse_result_to_parsed_document(_parse_result_with_tree())
        table_block = next(b for b in pd.blocks if b.type is ContentBlockType.TABLE)
        assert "| A | B |" in table_block.content

    def test_empty_tree_yields_empty_blocks(self):
        """ParseResult без дерева -> пустой blocks, пустой text не падает."""
        pr = ParseResult(
            document_id=uuid4().hex,
            filename="t.docx",
            format="dedoc",
            markdown="",
            tree=None,
            metadata={},
        )
        pd = map_parse_result_to_parsed_document(pr)
        assert pd.blocks == []
        assert pd.text == ""


# --------------------------------------------------------------------------- #
# Шаг 4: chunk_mapper.py
# --------------------------------------------------------------------------- #


class TestChunkMapper:
    """Внутренний Chunk -> архитектурный Chunk (BEAR)."""

    def test_id_is_valid_uuid(self):
        """chunk.id — валидный UUID (критерий приёмки Шага 7)."""
        ac = map_internal_to_architectural_chunk(_internal_chunk(), uuid4(), 1, 0)
        assert isinstance(ac.id, UUID)

    def test_text_is_text_for_embedding(self):
        """chunk.text = internal_chunk.text_for_embedding (именно это поле)."""
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(text_for_embedding="[crumb] обогащённый текст"),
            uuid4(),
            1,
            0,
        )
        assert ac.text == "[crumb] обогащённый текст"

    def test_text_not_empty(self):
        """chunk.text не пустой (критерий приёмки Шага 7)."""
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(text_for_embedding="текст"), uuid4(), 1, 0
        )
        assert ac.text

    def test_text_fallback_when_no_embedding(self):
        """text_for_embedding=None -> откат на internal_chunk.text (не пустой)."""
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(text="raw текст", text_for_embedding=None),
            uuid4(),
            1,
            0,
        )
        assert ac.text == "raw текст"

    def test_metadata_has_required_keys(self):
        """metadata содержит parent_id, block_type, breadcrumb (критерий Шага 7)."""
        ac = map_internal_to_architectural_chunk(_internal_chunk(), uuid4(), 1, 0)
        assert "parent_id" in ac.metadata
        assert "block_type" in ac.metadata
        assert "breadcrumb" in ac.metadata

    def test_parent_id_from_internal_chunk(self):
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(parent_id="h_5"), uuid4(), 1, 0
        )
        assert ac.metadata["parent_id"] == "h_5"

    def test_parent_id_fallback_to_doc_group(self):
        """parent_id=None -> f"doc_{document_id}" (бессемейные чанки)."""
        doc_id = uuid4()
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(parent_id=None), doc_id, 1, 0
        )
        assert ac.metadata["parent_id"] == f"doc_{doc_id}"

    def test_block_type_text(self):
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(chunk_type="text"), uuid4(), 1, 0
        )
        assert ac.metadata["block_type"] == "text"
        assert ac.metadata["table_representation"] is None

    def test_block_type_table(self):
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(chunk_type="table"), uuid4(), 1, 0
        )
        assert ac.metadata["block_type"] == "table"
        assert ac.metadata["table_representation"] == "markdown"

    def test_is_table_helper(self):
        assert is_table(_internal_chunk(chunk_type="table")) is True
        assert is_table(_internal_chunk(chunk_type="text")) is False
        assert is_table(_internal_chunk(chunk_type=None)) is False

    def test_token_count_positive(self):
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(text_for_embedding="осмысленный текст чанка"),
            uuid4(),
            1,
            0,
        )
        assert ac.token_count > 0

    def test_token_count_uses_existing_count_tokens(self):
        """token_count считается существующей count_tokens (tiktoken cl100k_base)."""
        from rag_parser.chunkers.base import count_tokens

        text = "[crumb] тестовый текст для проверки токенов"
        ac = map_internal_to_architectural_chunk(
            _internal_chunk(text_for_embedding=text), uuid4(), 1, 0
        )
        assert ac.token_count == count_tokens(text)

    def test_document_id_and_version_propagated(self):
        doc_id = uuid4()
        ac = map_internal_to_architectural_chunk(_internal_chunk(), doc_id, 7, 3)
        assert ac.document_id == doc_id
        assert ac.document_version == 7
        assert ac.sequence == 3

    def test_metadata_mime_type_is_none(self):
        ac = map_internal_to_architectural_chunk(_internal_chunk(), uuid4(), 1, 0)
        assert ac.metadata["mime_type"] is None

    def test_metadata_attachments_is_empty_list(self):
        ac = map_internal_to_architectural_chunk(_internal_chunk(), uuid4(), 1, 0)
        assert ac.metadata["attachments"] == []


# --------------------------------------------------------------------------- #
# Шаг 5: chunk_architectural (HybridRagChunker + BaseChunker default)
# --------------------------------------------------------------------------- #


class TestChunkArchitectural:
    """Метод chunk_architectural: ParsedDocument -> List[архитектурный Chunk]."""

    def test_hybrid_returns_architectural_chunks(self):
        from rag_parser.chunkers.hybrid import HybridRagChunker

        pd = ParsedDocument(
            document_id=uuid4(),
            text="# Заголовок\n\nТекст абзаца под заголовком.",
            blocks=[],
            attachments=[],
            metadata={},
        )
        chunker = HybridRagChunker(text_strategy="token")
        chunks = chunker.chunk_architectural(pd)
        assert len(chunks) >= 1
        for c in chunks:
            assert isinstance(c, ArchitecturalChunk)
            assert isinstance(c.id, UUID)
            assert c.text  # не пустой
            assert "parent_id" in c.metadata
            assert "block_type" in c.metadata
            assert "breadcrumb" in c.metadata

    def test_sequence_is_global_enumerate(self):
        """sequence = 0, 1, 2, ... по enumerate (глобальный, не внутри-родителя)."""
        from rag_parser.chunkers.hybrid import HybridRagChunker

        pd = ParsedDocument(
            document_id=uuid4(),
            text="# H1\n\nАбзац 1.\n\nАбзац 2.\n\nАбзац 3.",
            blocks=[],
            attachments=[],
            metadata={},
        )
        chunker = HybridRagChunker(text_strategy="token", chunk_size=50, overlap=0)
        chunks = chunker.chunk_architectural(pd)
        sequences = [c.sequence for c in chunks]
        assert sequences == list(range(len(chunks)))

    def test_document_id_propagated(self):
        from rag_parser.chunkers.hybrid import HybridRagChunker

        doc_id = uuid4()
        pd = ParsedDocument(
            document_id=doc_id,
            text="Простой текст.",
            blocks=[],
            attachments=[],
            metadata={},
        )
        chunker = HybridRagChunker(text_strategy="token")
        chunks = chunker.chunk_architectural(pd)
        assert all(c.document_id == doc_id for c in chunks)

    def test_token_chunker_inherits_default(self):
        """TokenChunkerRAG наследует chunk_architectural от BaseChunker."""
        from rag_parser.chunkers.token import TokenChunkerRAG

        pd = ParsedDocument(
            document_id=uuid4(),
            text="Текст для чанкинга токенами.",
            blocks=[],
            attachments=[],
            metadata={},
        )
        chunker = TokenChunkerRAG(chunk_size=50, overlap=0)
        chunks = chunker.chunk_architectural(pd)
        assert len(chunks) >= 1
        assert all(isinstance(c.id, UUID) for c in chunks)

    def test_old_chunk_signature_unchanged(self):
        """Старый .chunk(text, metadata) всё ещё работает (не сломан)."""
        from rag_parser.chunkers.token import TokenChunkerRAG

        chunker = TokenChunkerRAG(chunk_size=50, overlap=0)
        # старый вызов — без ParsedDocument, без document_version
        chunks = chunker.chunk("Текст для старого вызова.", {"source": "t"})
        assert len(chunks) >= 1
        # возвращает ВНУТРЕННИЕ чанки (не архитектурные)
        assert all(isinstance(c, InternalChunk) for c in chunks)

    def test_hybrid_old_chunk_with_tree_still_works(self):
        """HybridRagChunker.chunk(text, metadata, tree=) не сломан."""
        from rag_parser.chunkers.hybrid import HybridRagChunker

        pr = _parse_result_with_tree()
        chunker = HybridRagChunker(text_strategy="token")
        chunks = chunker.chunk(pr.markdown, {"source": "t"}, tree=pr.tree)
        assert len(chunks) >= 1
        assert all(isinstance(c, InternalChunk) for c in chunks)


# --------------------------------------------------------------------------- #
# Шаг 6: API — ArchitecturalChunkResponse + architectural_chunks
# --------------------------------------------------------------------------- #


class TestApiSchema:
    """Новая Pydantic-схема и поле ParseResponse.architectural_chunks."""

    def test_architectural_chunk_response_schema(self):
        from rag_parser.api.schemas import ArchitecturalChunkResponse

        resp = ArchitecturalChunkResponse(
            id=str(uuid4()),
            document_id=str(uuid4()),
            document_version=1,
            sequence=0,
            text="текст",
            token_count=2,
            metadata={"parent_id": "h_1", "block_type": "text"},
        )
        assert resp.id  # str
        assert resp.metadata["parent_id"] == "h_1"

    def test_parse_response_has_architectural_chunks_field(self):
        from rag_parser.api.schemas import ParseResponse

        # поле опциональное — default None
        fields = ParseResponse.model_fields
        assert "architectural_chunks" in fields
        # старое поле chunks на месте (не удалено)
        assert "chunks" in fields

    def test_parse_response_old_fields_intact(self):
        """Старые поля ParseResponse не удалены и не переименованы."""
        from rag_parser.api.schemas import ParseResponse

        fields = ParseResponse.model_fields
        for required in (
            "document_id",
            "filename",
            "format",
            "markdown",
            "chunks",
            "tree",
            "metadata",
        ):
            assert required in fields, f"старое поле {required} пропало"

    def test_endpoint_returns_architectural_chunks(self, tmp_path):
        """Сквозной тест: POST /api/parse отдаёт architectural_chunks."""
        import docx
        from fastapi.testclient import TestClient
        from rag_parser.api.routes import router
        from fastapi import FastAPI

        # минимальный DOCX
        document = docx.Document()
        document.add_paragraph("Тестовый абзац документа.")
        source = tmp_path / "arch.docx"
        document.save(source)

        app = FastAPI()
        app.include_router(router)
        client = TestClient(app)
        with open(source, "rb") as f:
            response = client.post(
                "/api/parse",
                files={"file": ("arch.docx", f, "application/octet-stream")},
                data={"chunking_type": "fixed", "chunk_size": "200", "overlap": "0"},
            )
        assert response.status_code == 200, response.text
        body = response.json()
        # старое поле на месте
        assert "chunks" in body
        # новое поле появилось и заполнено
        assert "architectural_chunks" in body
        arch = body["architectural_chunks"]
        assert arch is not None
        assert len(arch) >= 1
        for c in arch:
            # id — валидный UUID-строка
            UUID(c["id"])
            UUID(c["document_id"])
            assert c["text"]  # не пустой
            assert "parent_id" in c["metadata"]
            assert "block_type" in c["metadata"]
            assert "breadcrumb" in c["metadata"]
            assert c["sequence"] >= 0
