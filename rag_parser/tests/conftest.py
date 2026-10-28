from pathlib import Path

import pytest

TXT_UTF8 = "Первая строка документа.\nВторая строка документа.\nТретья строка.\n"
TXT_LATIN1 = "Café résumé naïve\nSecond line\n"
MD_DOC = (
    "# Заголовок первого уровня\n\n"
    "Абзац с текстом и **жирными** словами.\n\n"
    "## Подзаголовок\n\n"
    "- пункт один\n- пункт два\n\n"
    "| Колонка | Значение |\n|---|---|\n| A | 1 |\n"
)
CSV_SEMICOLON = "Колонка1;Колонка2;Колонка3\nЗначение1;Значение2;Значение3\n"
CSV_COMMA = "name,qty,price\napple,3,1.5\nbanana,10,0.75\n"
HTML_DOC = (
    "<!DOCTYPE html><html><head><title>T</title>"
    "<style>body{color:red}</style><script>var x=1;</script></head>"
    "<body><h1>Заголовок</h1><p>Текст со <a href='https://example.com'>ссылкой</a>.</p>"
    "<h2>Подзаголовок</h2><table><tr><th>A</th><th>B</th></tr>"
    "<tr><td>1</td><td>2</td></tr></table></body></html>"
)
HTML_MINIMAL = "<html><body><h3>Title</h3><ul><li>one</li><li>two</li></ul></body></html>"
SENTENCE_TEXT = (
    "Машинное обучение изучает алгоритмы, которые учатся на данных. "
    "Модели достигают высокой точности на проверочных наборах. "
    "Классические алгоритмы остаются интерпретируемыми и быстрыми. "
    "Глубокое обучение требует больших объёмов размеченных данных. "
    "Современные системы достигают уровня человека в ряде задач. "
) * 8


def write_bytes(path: Path, data: bytes) -> Path:
    path.write_bytes(data)
    return path


def make_text_files(tmp_path: Path) -> dict:
    return {
        "txt_utf8": write_bytes(tmp_path / "utf8.txt", TXT_UTF8.encode("utf-8")),
        "txt_latin1": write_bytes(tmp_path / "latin1.txt", TXT_LATIN1.encode("latin-1")),
        "md": write_bytes(tmp_path / "doc.md", MD_DOC.encode("utf-8")),
    }


def make_csv_files(tmp_path: Path) -> dict:
    return {
        "csv_semicolon": write_bytes(
            tmp_path / "semicolon.csv", CSV_SEMICOLON.encode("utf-8")
        ),
        "csv_comma": write_bytes(tmp_path / "comma.csv", CSV_COMMA.encode("utf-8")),
        "tsv": write_bytes(
            tmp_path / "data.tsv",
            "name\tqty\tprice\napple\t3\t1.5\nbanana\t10\t0.75\n".encode("utf-8"),
        ),
    }


def make_html_files(tmp_path: Path) -> dict:
    return {
        "html_full": write_bytes(tmp_path / "full.html", HTML_DOC.encode("utf-8")),
        "html_minimal": write_bytes(
            tmp_path / "minimal.html", HTML_MINIMAL.encode("utf-8")
        ),
        "htm": write_bytes(tmp_path / "page.htm", HTML_DOC.encode("utf-8")),
    }


COMPLEX_DOCX = (
    Path(__file__).resolve().parent.parent / "test_data" / "complex_structured.docx"
)


@pytest.fixture(scope="session")
def complex_docx() -> Path:
    """Path to the structured DOCX fixture, generated on demand."""
    if not COMPLEX_DOCX.exists():
        from rag_parser.scripts.generate_test_doc import generate

        generate(COMPLEX_DOCX)
    return COMPLEX_DOCX


def make_docx_files(tmp_path: Path) -> dict:
    from docx import Document

    first = Document()
    first.add_heading("Заголовок документа", level=1)
    first.add_paragraph("Первый абзац текста.")
    first.add_heading("Второй раздел", level=2)
    first.add_paragraph("Второй абзац текста.")
    first_path = tmp_path / "first.docx"
    first.save(first_path)

    second = Document()
    second.add_paragraph("Абзац с таблицей.")
    table = second.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    second_path = tmp_path / "second.docx"
    second.save(second_path)
    return {"docx_headings": first_path, "docx_table": second_path}


def make_xlsx_files(tmp_path: Path) -> dict:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Данные"
    sheet.append(["Колонка1", "Колонка2"])
    sheet.append(["Значение1", 2])
    extra = workbook.create_sheet("Итоги")
    extra.append(["Показатель", "Значение"])
    extra.append(["Выручка", 100])
    first_path = tmp_path / "multi.xlsx"
    workbook.save(first_path)

    single = Workbook()
    only = single.active
    only.title = "Single"
    only.append(["id", "name"])
    for index in range(5):
        only.append([index, f"row-{index}"])
    second_path = tmp_path / "single.xlsx"
    single.save(second_path)

    merged = Workbook()
    sheet = merged.active
    sheet.title = "Data"
    sheet.append(["Report Title", None, None])
    sheet.merge_cells("A1:C1")
    sheet.append(["H1", "H2", "H3"])
    sheet.append(["a", "b", "c"])
    sheet.append(["merged-row", None, None])
    sheet.merge_cells("A4:C4")
    sheet.append(["d", "e", "f"])
    sheet.append(["2D-block", "B6", "C6"])
    sheet.merge_cells("B6:C7")
    sheet.append(["A7", None, None])
    vertical = merged.create_sheet("Vertical")
    vertical.append(["top", "x"])
    vertical.merge_cells("A1:A3")
    vertical.append([None, "y"])
    vertical.append([None, "z"])
    # Анти-регрессия лево-право vs верх-вниз: мёрдж B2:B5 во 2-й колонке,
    # в 1-й колонке каждой строки — РАЗЛИЧНЫЕ значения.
    vcol = merged.create_sheet("VCol")
    vcol.append(["id", "Группа"])
    vcol.append(["r1", "TOP"])
    vcol.merge_cells("B2:B5")
    vcol.append(["r2", None])
    vcol.append(["r3", None])
    vcol.append(["r4", None])
    merged_path = tmp_path / "merged.xlsx"
    merged.save(merged_path)
    return {
        "xlsx_multi": first_path,
        "xlsx_single": second_path,
        "xlsx_merged": merged_path,
    }


def make_xls_files(tmp_path: Path) -> dict:
    """Копирует предсгенерированный .xls (xlwt — только локально для генерации)."""
    import shutil

    source = Path(__file__).resolve().parent.parent / "test_data" / "sample.xls"
    target = tmp_path / "sample.xls"
    shutil.copy(source, target)
    return {"xls": target}


def make_header_docs(tmp_path: Path) -> dict:
    """Пара фикстур для HeaderPromoter: со стилями Word и без (визуальные)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    styled = Document()
    styled.add_heading("Раздел 1. Общие сведения", level=1)
    styled.add_paragraph("Первый абзац со сведениями о документе.")
    styled.add_heading("1.1 Метрики качества", level=2)
    styled.add_paragraph("Абзац про метрики качества чанкинга.")
    styled_path = tmp_path / "styled.docx"
    styled.save(styled_path)

    unstyled = Document()
    title = unstyled.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Раздел 2. Основные показатели")
    run.bold = True
    run.font.size = Pt(16)
    unstyled.add_paragraph(
        "Длинный абзац тела документа, в котором подробно описываются "
        "основные показатели работы системы в отчётном периоде. " * 4
    )
    subtitle = unstyled.add_paragraph()
    run2 = subtitle.add_run("2.1 Детализация по кварталам")
    run2.bold = True
    run2.font.size = Pt(14)
    unstyled.add_paragraph("Короткий абзац детализации.")
    unstyled_path = tmp_path / "unstyled.docx"
    unstyled.save(unstyled_path)
    return {"styled": styled_path, "unstyled": unstyled_path}


def make_annual_report_doc(tmp_path: Path) -> Path:
    """Фикстура в стиле 2023.docx: TOC + тело с N)-пунктом и подуровнями."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Годовой отчет акционерного общества за 2023 год.")
    for index in range(1, 7):
        doc.add_paragraph(f"Раздел {index}. Строка оглавления годового отчета")
    doc.add_paragraph("Раздел 1. Информация о годовом отчете общества")
    doc.add_paragraph("Тело первого раздела с общими сведениями.")
    doc.add_paragraph("1) достоверность годового отчета подтверждена ревизией")
    doc.add_paragraph("2) годовой отчет предварительно утвержден советом директоров")
    doc.add_paragraph("3) годовой отчет утвержден общим собранием акционеров")
    doc.add_paragraph("Раздел 2. Сведения об обществе и его деятельности")
    doc.add_paragraph("2.1. Полное наименование общества.")
    doc.add_paragraph("Тело подраздела с реквизитами общества.")
    path = tmp_path / "annual_2023_style.docx"
    doc.save(path)
    return path


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _text_ops(lines, size=12):
    ops = []
    for index, line in enumerate(lines):
        if line:
            ops.append(
                f"BT /F1 {size} Tf 72 {720 - index * 20} Td ({_pdf_escape(line)}) Tj ET"
            )
    return ops


def _table_ops(table, x0=72, y0=400, cell_width=160, row_height=30):
    rows, columns = len(table), len(table[0])
    ops = ["0.5 w"]
    for row in range(rows + 1):
        ops.append(
            f"{x0} {y0 + row * row_height} m {x0 + columns * cell_width} {y0 + row * row_height} l S"
        )
    for column in range(columns + 1):
        ops.append(
            f"{x0 + column * cell_width} {y0} m {x0 + column * cell_width} {y0 + rows * row_height} l S"
        )
    for row, values in enumerate(table):
        for column, value in enumerate(values):
            position = y0 + rows * row_height - row * row_height - 18
            ops.append(
                f"BT /F1 11 Tf {x0 + column * cell_width + 8} {position} Td "
                f"({_pdf_escape(str(value))}) Tj ET"
            )
    return ops


def _build_pdf(pages, tables=None, path=None):
    """Builds a small valid PDF without external dependencies.

    The embedded base-14 Helvetica font covers latin text only, so PDF
    fixtures use ascii content.
    """
    operations = [list(page) for page in pages]
    for table in tables or []:
        operations.append(_table_ops(table))
    if not operations:
        operations = [[]]

    page_count = len(operations)
    first_content = 3 + page_count
    font_object = 3 + 2 * page_count
    objects = {}
    kids = " ".join(f"{3 + index} 0 R" for index in range(page_count))
    objects[1] = "<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>"
    for index, ops in enumerate(operations):
        objects[3 + index] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {first_content + index} 0 R "
            f"/Resources << /Font << /F1 {font_object} 0 R >> >> >>"
        )
        stream = "\n".join(ops)
        objects[first_content + index] = (
            f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"
        )
    objects[font_object] = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n{objects[number]}\nendobj\n".encode("latin-1")
    xref_offset = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for number in range(1, size):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    return write_bytes(path, bytes(out))


def pdf_page(title, body=""):
    return _text_ops([title, *([body] if body else [])])


def make_pdf_files(tmp_path: Path) -> dict:
    mixed_pages = [
        pdf_page("Alpha page contains plenty of words to classify confidently."),
        # Текстовая страница с таблицей (таблица в содержимом страницы).
        list(pdf_page("Quarterly data page with a table below."))
        + _table_ops([["Region", "Q1"], ["North", "100"], ["South", "90"]]),
        # Страница-картинка (нет текстового слоя) внутри текстового PDF.
        [],
        # Страница с мусорным текстовым слоем (битые CMap-маркеры).
        _text_ops(["(cid:12)(cid:3)(cid:44)(cid:7)"] * 12, size=9),
    ]
    return {
        "pdf_text": _build_pdf(
            [pdf_page("First page of the document."), pdf_page("Second page body.")],
            None,
            tmp_path / "text.pdf",
        ),
        "pdf_table": _build_pdf(
            [pdf_page("Quarterly report")],
            [[["Region", "Q1", "Q2"], ["North", "100", "120"], ["South", "90", "95"]]],
            tmp_path / "table.pdf",
        ),
        # Валидный PDF БЕЗ текстового слоя — эмуляция скана: страница есть,
        # текстовых операторов нет. Обрабатывается через OCR-ветку.
        "pdf_scan": _build_pdf(
            [[]],  # ни одного BT/Tj — только пустая страница
            None,
            tmp_path / "scan.pdf",
        ),
        # Смешанный PDF: текст → текст+таблица → скан → мусорный слой.
        "pdf_mixed": _build_pdf(mixed_pages, None, tmp_path / "mixed.pdf"),
    }


def make_invalid_files(tmp_path: Path) -> dict:
    return {
        "broken_pdf": write_bytes(tmp_path / "broken.pdf", b"%PDF-1.4\nnot really a pdf"),
        "empty_pdf": write_bytes(tmp_path / "empty.pdf", b""),
        "not_a_pdf": write_bytes(tmp_path / "fake.pdf", b"plain text file"),
        "empty_txt": write_bytes(tmp_path / "empty.txt", b""),
        "broken_docx": write_bytes(tmp_path / "broken.docx", b"not a zip container"),
        "broken_xlsx": write_bytes(tmp_path / "broken.xlsx", b"not a zip container"),
        "empty_csv": write_bytes(tmp_path / "empty.csv", b"\n"),
        "blank_csv": write_bytes(tmp_path / "blank.csv", b"   \n"),
        "empty_html": write_bytes(tmp_path / "empty.html", b""),
    }


@pytest.fixture
def text_files(tmp_path):
    return make_text_files(tmp_path)


@pytest.fixture
def csv_files(tmp_path):
    return make_csv_files(tmp_path)


@pytest.fixture
def html_files(tmp_path):
    return make_html_files(tmp_path)


@pytest.fixture
def docx_files(tmp_path):
    return make_docx_files(tmp_path)


@pytest.fixture
def xlsx_files(tmp_path):
    return make_xlsx_files(tmp_path)


@pytest.fixture
def xls_files(tmp_path):
    return make_xls_files(tmp_path)


@pytest.fixture
def header_docs(tmp_path):
    return make_header_docs(tmp_path)


@pytest.fixture
def annual_report_doc(tmp_path):
    return make_annual_report_doc(tmp_path)


@pytest.fixture
def pdf_files(tmp_path):
    return make_pdf_files(tmp_path)


@pytest.fixture
def invalid_files(tmp_path):
    return make_invalid_files(tmp_path)


@pytest.fixture
def sample_text():
    return SENTENCE_TEXT


@pytest.fixture
def client(tmp_path, monkeypatch):
    """TestClient приложения с изолированным document_store.

    Общий фикстур: используется и в test_api.py, и в test_retrieval.py
    (интеграционные тесты Child-Parent эндпоинта).
    """
    from fastapi.testclient import TestClient

    from rag_parser.api.dependencies import DocumentStore
    from rag_parser.core.config import settings

    monkeypatch.setattr(
        settings, "DEDOC_OUTPUT_DIR", str(tmp_path / "output" / "dedoc_results")
    )
    # Мокаем PDF-пайплайн для старых API-тестов: Docling не установлен в
    # тестовом окружении (промпт: «в юнит-тестах Docling не загружается»).
    # ``test_pdf_pipeline.py`` мокает Docling сам и не использует фикстуру
    # ``client``, поэтому autouse-мок не нужен.
    from rag_parser.parsers import pdf_pipeline
    from rag_parser.parsers.pdf_pipeline import PdfPipelineResult

    def _mock_run_pipeline(path, **kwargs):
        # Битый PDF (сигнатура %PDF есть, но содержимое мусор) → ошибка,
        # чтобы тесты ``test_broken_pdf`` получали 400, а не 200.
        import fitz
        from rag_parser.core.interfaces import DocumentParseError
        try:
            doc = fitz.open(str(path))
            total = len(doc)
            doc.close()
        except Exception as exc:
            raise DocumentParseError(
                f"PDF pipeline mock: cannot open {path.name}: {exc}",
                filename=path.name,
            ) from exc
        return PdfPipelineResult(
            markdown=f"PDF pipeline mocked: {path.name}",
            metadata={"source_pipeline": "pdf_router", "pdf_type": "text",
                      "total_pages": total, "text_pages": list(range(total)),
                      "scan_pages": [],
                      "routing": {"text_parser_pages": list(range(total)),
                                  "docling_pages": [], "docling_ocr_pages": [],
                                  "table_runs": []},
                      "docling_runs": [], "lost_tables": [], "timings": {},
                      "format": "dedoc", "extension": ".pdf",
                      "size_bytes": path.stat().st_size if path.exists() else 0,
                      "parser": "pdf_pipeline"},
        )

    monkeypatch.setattr(pdf_pipeline, "run_pipeline", _mock_run_pipeline)

    import rag_parser.api.routes as routes
    from rag_parser.main import app

    previous = routes.document_store
    routes.document_store = DocumentStore(max_items=16)
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client
    routes.document_store = previous
