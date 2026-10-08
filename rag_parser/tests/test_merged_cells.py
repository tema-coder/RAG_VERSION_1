"""Тесты объединённых ячеек: точная заливка по диапазонам + эвристика."""

import logging

import pytest

from rag_parser.parsers.dedoc_parser import DedocParser
from rag_parser.parsers.merged_cells import (
    fill_by_ranges,
    fill_document_tables,
    fill_workbook_tables,
    filled_table_markdown,
    forward_fill_matrix,
    normalize_ranges,
    sheet_merged_ranges,
)


class TestForwardFillMatrix:
    def test_horizontal(self):
        assert forward_fill_matrix([["T", "", ""], ["a", "b", "c"]]) == [
            ["T", "T", "T"],
            ["a", "b", "c"],
        ]

    def test_vertical(self):
        assert forward_fill_matrix([["top", "x"], ["", "y"], ["", "z"]]) == [
            ["top", "x"],
            ["top", "y"],
            ["top", "z"],
        ]

    def test_2d_block(self):
        # Сначала заливка слева, затем сверху (как в старом xlsx-парсере).
        assert forward_fill_matrix([["A", "B", "C"], ["D", "", ""]]) == [
            ["A", "B", "C"],
            ["D", "D", "D"],
        ]
        assert forward_fill_matrix([["A", "B"], ["", ""]]) == [
            ["A", "B"],
            ["A", "B"],
        ]

    def test_empty_matrix(self):
        assert forward_fill_matrix([]) == []


class TestFillDocumentTables:
    def _table(self):
        return {
            "cells": [
                [{"lines": [{"text": "Report Title"}]}, {"lines": [{"text": ""}]},
                 {"lines": [{"text": ""}]}],
                [{"lines": [{"text": "a"}]}, {"lines": [{"text": "b"}]},
                 {"lines": [{"text": "c"}]}],
            ]
        }

    def test_dict_table(self):
        table = self._table()
        assert fill_document_tables([table]) == 2
        assert table["cells"][0][1]["lines"][0]["text"] == "Report Title"
        assert table["cells"][0][2]["lines"][0]["text"] == "Report Title"

    def test_line_key_and_missing_lines(self):
        table = {"cells": [[{"lines": [{"line": "T"}]}, {"lines": []},
                            {"other": True}]]}
        assert fill_document_tables([table]) == 2
        assert table["cells"][0][1]["lines"][0]["text"] == "T"
        assert table["cells"][0][2]["lines"][0]["text"] == "T"

    def test_object_cells(self):
        from rag_parser.parsers.dedoc_tree import extract_line_text

        class ObjLine:
            def __init__(self, text):
                self._text = text

            @property
            def line(self):
                return self._text

            def set_line(self, value):
                self._text = value

        class ObjCell:
            def __init__(self, text):
                self.lines = [ObjLine(text)]
                self.colspan = 1
                self.rowspan = 1
                self.invisible = False

        class ObjTable:
            def __init__(self, cells):
                self.cells = cells

        table = ObjTable([[ObjCell("T"), ObjCell("")], [ObjCell(""), ObjCell("x")]])
        assert fill_document_tables([table]) == 2
        assert extract_line_text(table.cells[0][1].lines[0]) == "T"
        assert extract_line_text(table.cells[1][0].lines[0]) == "T"

    def test_no_tables(self):
        assert fill_document_tables([]) == 0
        assert fill_document_tables(None) == 0

    def test_garbage_cells_ignored(self):
        assert fill_document_tables([{"cells": "nope"}]) == 0
        assert fill_document_tables([{"no_cells": []}]) == 0
        assert fill_document_tables([object()]) == 0

    def test_filled_markdown(self):
        markdown = filled_table_markdown(self._table())
        assert markdown.splitlines()[0].count("Report Title") == 3


class TestMergedCellsEndToEnd:
    def test_parser_fills_merged(self, tmp_path, monkeypatch):
        from rag_parser.core.config import settings

        monkeypatch.setattr(
            settings, "DEDOC_OUTPUT_DIR", str(tmp_path / "output" / "dedoc_results")
        )
        fixture = (
            __import__("pathlib").Path(__file__).resolve().parent.parent
            / "test_data"
            / "merged_cells.xlsx"
        )
        result = DedocParser().parse(fixture)
        lines = result.markdown.splitlines()
        title_rows = [line for line in lines if "Report Title" in line]
        assert title_rows
        assert title_rows[0].count("Report Title") == 3
        merged_rows = [line for line in lines if "merged-row" in line]
        assert merged_rows
        assert merged_rows[0].count("merged-row") == 3
        # вертикальный мёрдж: значение продублировано вниз
        top_rows = [line for line in lines if "top" in line]
        assert len(top_rows) >= 3

    def test_tmp_merged_fixture(self, xlsx_files, tmp_path, monkeypatch):
        from rag_parser.core.config import settings

        monkeypatch.setattr(
            settings, "DEDOC_OUTPUT_DIR", str(tmp_path / "output" / "dedoc_results")
        )
        result = DedocParser().parse(xlsx_files["xlsx_merged"])
        assert "Report Title" in result.markdown
        title_rows = [
            line for line in result.markdown.splitlines() if "Report Title" in line
        ]
        assert title_rows[0].count("Report Title") == 3
        two_d = [line for line in result.markdown.splitlines() if "B6" in line]
        assert two_d  # 2D-диапазон B6:C7 заполнен

    def test_vertical_merge_second_column_top_down(self, xlsx_files, tmp_path,
                                                  monkeypatch):
        """Анти-регрессия: мёрдж B2:B5 заливается СВЕРХУ (TOP), не слева."""
        from rag_parser.core.config import settings

        monkeypatch.setattr(
            settings, "DEDOC_OUTPUT_DIR", str(tmp_path / "output" / "dedoc_results")
        )
        result = DedocParser().parse(xlsx_files["xlsx_merged"])
        vcol_rows = [
            [cell.strip() for cell in line.strip().strip("|").split("|")]
            for line in result.markdown.splitlines()
            if line.strip().startswith("| r")
        ]
        assert len(vcol_rows) == 4
        for row in vcol_rows:
            assert row[1] == "TOP", row  # не r1/r2/r3/r4 из колонки 1
        assert [row[0] for row in vcol_rows] == ["r1", "r2", "r3", "r4"]


def _dict_table(matrix):
    return {"cells": [[{"lines": [{"text": value}]} for value in row]
                      for row in matrix]}


def _texts(table):
    from rag_parser.parsers.dedoc_tree import cell_to_text

    return [[cell_to_text(cell) for cell in row] for row in table["cells"]]


class TestFillByRanges:
    def test_vertical_in_second_column(self):
        table = _dict_table([["id", "Группа"], ["r1", "TOP"], ["r2", ""],
                             ["r3", ""], ["r4", ""]])
        assert fill_by_ranges(table, [(1, 5, 1, 2)]) == 3
        assert _texts(table) == [["id", "Группа"], ["r1", "TOP"], ["r2", "TOP"],
                                 ["r3", "TOP"], ["r4", "TOP"]]

    def test_horizontal_and_2d(self):
        table = _dict_table([["T", "", ""], ["a", "b", "c"]])
        assert fill_by_ranges(table, [(0, 1, 0, 3)]) == 2
        assert _texts(table)[0] == ["T", "T", "T"]

    def test_empty_top_left_skipped(self):
        table = _dict_table([["", ""], ["", "x"]])
        assert fill_by_ranges(table, [(0, 2, 0, 1)]) == 0

    def test_normalize_xlrd_and_openpyxl(self):
        assert normalize_ranges([(1, 4, 1, 2)]) == [(1, 4, 1, 2)]

        class FakeRange:
            bounds = (2, 2, 2, 5)  # B2:B5 1-based inclusive

        assert normalize_ranges([FakeRange()]) == [(1, 5, 1, 2)]
        assert normalize_ranges([object()]) == []
        assert normalize_ranges(None) == []


class TestFillWorkbookTables:
    def _workbook(self, path):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "VCol"
        sheet.append(["id", "Группа"])
        sheet.append(["r1", "TOP"])
        sheet.merge_cells("B2:B5")
        sheet.append(["r2", None])
        sheet.append(["r3", None])
        sheet.append(["r4", None])
        workbook.save(path)
        return path

    def test_precise_by_ranges(self, tmp_path):
        path = self._workbook(tmp_path / "v.xlsx")
        assert sheet_merged_ranges(path) == [[(1, 5, 1, 2)]]
        table = _dict_table([["id", "Группа"], ["r1", "TOP"], ["r2", ""],
                             ["r3", ""], ["r4", ""]])
        assert fill_workbook_tables([table], path) == 3
        assert [row[1] for row in _texts(table)[1:]] == ["TOP"] * 4

    def test_dimension_mismatch_falls_back(self, tmp_path, caplog):
        path = self._workbook(tmp_path / "v.xlsx")
        # Таблица 2x2 против листа 5x2: сверка обязана сработать, не вслепую.
        table = _dict_table([["a", ""], ["", "b"]])
        with caplog.at_level(logging.WARNING):
            filled = fill_workbook_tables([table], path)
        assert filled == 2  # эвристика: (0,1)->a, (1,0)->a
        assert _texts(table) == [["a", "a"], ["a", "b"]]
        assert any("mismatch" in record.message for record in caplog.records)

    def test_missing_file_falls_back(self, tmp_path, caplog):
        table = _dict_table([["T", ""], ["a", "b"]])
        with caplog.at_level(logging.WARNING):
            assert fill_workbook_tables([table], tmp_path / "nope.xlsx") == 1
        assert _texts(table)[0] == ["T", "T"]

    def test_no_tables(self, tmp_path):
        assert fill_workbook_tables([], tmp_path / "whatever.xlsx") == 0


class TestXlsMergeTransfer:
    def test_merged_cells_survive_conversion(self, tmp_path):
        import xlrd
        from openpyxl import load_workbook

        from rag_parser.parsers.dedoc_parser import convert_xls_to_xlsx

        try:
            import xlwt
        except ImportError:
            pytest.skip("xlwt доступен только в dev-окружении")

        source = tmp_path / "m.xls"
        book = xlwt.Workbook()
        sheet = book.add_sheet("VCol")
        sheet.write(0, 0, "id")
        sheet.write(0, 1, "Группа")
        sheet.write(1, 0, "r1")
        sheet.write_merge(1, 4, 1, 1, "TOP")
        sheet.write(2, 0, "r2")
        sheet.write(3, 0, "r3")
        sheet.write(4, 0, "r4")
        book.save(str(source))

        probe = xlrd.open_workbook(str(source), formatting_info=True)
        assert probe.sheet_by_index(0).merged_cells == [(1, 5, 1, 2)]

        converted = convert_xls_to_xlsx(source)
        try:
            assert sheet_merged_ranges(converted) == [[(1, 5, 1, 2)]]
            workbook = load_workbook(str(converted))
            try:
                assert workbook.worksheets[0].max_row == 5
            finally:
                workbook.close()
        finally:
            converted.unlink(missing_ok=True)
