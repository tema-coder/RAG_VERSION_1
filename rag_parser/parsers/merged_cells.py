"""Forward-fill объединённых ячеек xlsx/xls, спарсенных dedoc.

Проверено на реальном выводе dedoc (``ExcelReader``): ячейки объединённых
диапазонов приходят пустыми (``colspan == rowspan == 1``, текст ``""``),
значение хранится только в левой/верхней ячейке диапазона.

Два режима:

* точный — :func:`fill_by_ranges` по реальным диапазонам книги
  (``openpyxl`` для xlsx; при xlrd-конвертации xls мёрджи переносятся —
  см. ``DedocParser``). Значение левой-верхней ячейки копируется во все
  клетки диапазона, направление (вправо/вниз/2D) значения не имеет;
* эвристический — :func:`forward_fill_matrix` (слева, затем сверху),
  только как fallback без доступа к файлу. Известно неточен на
  вертикальных мёрджах с непустым левым соседом (невозможно различить
  без метаданных диапазона).

Работает как со словарями API-схемы (``model_dump``: ячейка
``{"lines": [{"text": ...}], ...}``), так и с объектами dedoc
(``CellWithMeta`` / ``Table``).
"""

import logging
from copy import deepcopy
from pathlib import Path
from typing import Any, List, Optional, Tuple

from rag_parser.parsers.dedoc_tree import cell_to_text

logger = logging.getLogger(__name__)

Range = Tuple[int, int, int, int]  # (rlo, rhi_excl, clo, chi_excl), 0-based


def is_empty_value(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def forward_fill_matrix(rows: List[List[str]]) -> List[List[str]]:
    """Заполняет пустые ячейки: сначала слева, затем сверху."""
    filled = [list(row) for row in rows]
    for row in filled:
        for index in range(1, len(row)):
            if is_empty_value(row[index]) and not is_empty_value(row[index - 1]):
                row[index] = row[index - 1]
    for row_index in range(1, len(filled)):
        above, current = filled[row_index - 1], filled[row_index]
        for index in range(len(current)):
            if is_empty_value(current[index]) and index < len(above):
                if not is_empty_value(above[index]):
                    current[index] = above[index]
    return filled


def _iter_cells(table: Any):
    if isinstance(table, dict):
        cells = table.get("cells")
    else:
        cells = getattr(table, "cells", None)
    if not isinstance(cells, list):
        return
    for row in cells:
        if not isinstance(row, (list, tuple)):
            continue
        for cell in row:
            yield cell


def _set_cell_text(cell: Any, text: str) -> None:
    if isinstance(cell, dict):
        lines = cell.get("lines")
        if isinstance(lines, list) and lines:
            first = lines[0]
            if isinstance(first, dict):
                if "text" in first:
                    first["text"] = text
                elif "line" in first:
                    first["line"] = text
                else:
                    lines[0] = {"text": text, "annotations": []}
            else:
                lines[0] = {"text": text, "annotations": []}
        else:
            cell["lines"] = [{"text": text, "annotations": []}]
        return
    lines = getattr(cell, "lines", None)
    if isinstance(lines, list) and lines:
        first = lines[0]
        if isinstance(first, dict):
            first["text"] = text
            return
        setter = getattr(first, "set_line", None)
        if callable(setter):
            try:
                setter(text)
                return
            except Exception:  # noqa: BLE001 - defensive
                pass
        try:
            first.line = text  # type: ignore[attr-defined]
            return
        except Exception:  # noqa: BLE001 - defensive
            pass
        try:
            first.text = text  # type: ignore[attr-defined]
            return
        except Exception:  # noqa: BLE001 - defensive
            pass
    try:
        from dedoc.data_structures import LineWithMeta

        cell.lines = [LineWithMeta(line=text)]
    except Exception:  # noqa: BLE001 - dedoc may be unavailable in tests
        try:
            cell.lines = [{"text": text, "annotations": []}]
        except Exception:  # noqa: BLE001 - defensive
            pass


def fill_merged_cells(table: Any) -> int:
    """Forward-fill одной таблицы dedoc на месте. Возвращает число заливок."""
    cells = list(_iter_cells(table))
    if not cells:
        return 0
    if isinstance(table, dict):
        matrix = table.get("cells")
    else:
        matrix = getattr(table, "cells", None)
    if not isinstance(matrix, list):
        return 0
    original = [
        [cell_to_text(cell) for cell in row]
        for row in matrix
        if isinstance(row, (list, tuple))
    ]
    filled = forward_fill_matrix(original)
    count = 0
    row_idx = 0
    for row in matrix:
        if not isinstance(row, (list, tuple)):
            continue
        for col_idx, cell in enumerate(row):
            if col_idx >= len(filled[row_idx]):
                continue
            if is_empty_value(cell_to_text(cell)) and not is_empty_value(
                filled[row_idx][col_idx]
            ):
                _set_cell_text(cell, filled[row_idx][col_idx])
                count += 1
        row_idx += 1
    return count


def fill_document_tables(tables: List[Any]) -> int:
    """Forward-fill всех таблиц документа. Возвращает суммарное число заливок."""
    return sum(fill_merged_cells(table) for table in tables or [])


def _table_matrix(table: Any) -> List[List[Any]]:
    if isinstance(table, dict):
        matrix = table.get("cells")
    else:
        matrix = getattr(table, "cells", None)
    if not isinstance(matrix, list):
        return []
    return [row for row in matrix if isinstance(row, (list, tuple))]


def normalize_ranges(ranges: List[Any]) -> List[Range]:
    """Диапазоны в единый вид (0-based, конец исключён).

    Принимает кортежи xlrd ``(rlo, rhi, clo, chi)`` и объекты
    ``openpyxl`` (``MergedCellRange`` с ``.bounds`` 1-based inclusive).
    """
    normalized: List[Range] = []
    for item in ranges or []:
        if isinstance(item, (list, tuple)) and len(item) == 4:
            rlo, rhi, clo, chi = (int(value) for value in item)
        else:
            bounds = getattr(item, "bounds", None)
            if bounds is None:
                continue
            min_col, min_row, max_col, max_row = bounds
            rlo, rhi, clo, chi = min_row - 1, max_row, min_col - 1, max_col
        if rhi > rlo and chi > clo and rlo >= 0 and clo >= 0:
            normalized.append((rlo, rhi, clo, chi))
    return normalized


def fill_by_ranges(table: Any, merged: List[Any]) -> int:
    """Точная заливка по диапазонам мёрджей. Возвращает число заливок."""
    matrix = _table_matrix(table)
    if not matrix:
        return 0
    count = 0
    for rlo, rhi, clo, chi in normalize_ranges(merged):
        if rlo >= len(matrix):
            continue
        top_row = matrix[rlo]
        if clo >= len(top_row):
            continue
        value = cell_to_text(top_row[clo])
        if is_empty_value(value):
            continue
        for row_idx in range(rlo, min(rhi, len(matrix))):
            row = matrix[row_idx]
            for col_idx in range(clo, min(chi, len(row))):
                if row_idx == rlo and col_idx == clo:
                    continue
                if is_empty_value(cell_to_text(row[col_idx])):
                    _set_cell_text(row[col_idx], value)
                    count += 1
    return count


def _workbook_layout(
    xlsx_path: Path,
) -> Optional[List[Tuple[Tuple[Optional[int], Optional[int]], List[Range]]]]:
    """``[((rows, cols), ranges)]`` по листам книги. None при ошибке."""
    try:
        from openpyxl import load_workbook

        # NB: без read_only — у ReadOnlyWorksheet нет .merged_cells.
        workbook = load_workbook(str(xlsx_path), data_only=True)
        try:
            return [((sheet.max_row, sheet.max_column),
                     normalize_ranges(list(sheet.merged_cells.ranges)))
                    for sheet in workbook.worksheets]
        finally:
            workbook.close()
    except Exception as exc:  # noqa: BLE001 - поврежденный файл и т.п.
        logger.warning("merged_ranges_unavailable", extra={"error": str(exc)})
        return None


def sheet_merged_ranges(xlsx_path: Path) -> Optional[List[List[Range]]]:
    """Диапазоны мёрджей по листам книги (порядок листов). None при ошибке."""
    layout = _workbook_layout(Path(xlsx_path))
    if layout is None:
        return None
    return [ranges for _, ranges in layout]


def fill_workbook_tables(tables: List[Any], xlsx_path: Path) -> int:
    """Заливка таблиц dedoc по диапазонам книги с потабличным fallback.

    Таблица ``i`` сопоставляется листу ``i`` (dedoc ``ExcelReader`` отдаёт
    строго по одной таблице на лист в порядке листов). Размерность
    сверяется: при несовпадении таблица заливается эвристикой
    (:func:`fill_merged_cells`) с warning — молча уехавшая не в ту таблицу
    заливка хуже эвристики.
    """
    if not tables:
        return 0
    layout = _workbook_layout(Path(xlsx_path))
    if layout is None or len(layout) != len(tables):
        logger.warning(
            "merged_sheet_mismatch",
            extra={"sheets": len(layout) if layout else 0,
                   "tables": len(tables)},
        )
        return fill_document_tables(tables)
    total = 0
    for index, (table, (dims, ranges)) in enumerate(zip(tables, layout)):
        matrix = _table_matrix(table)
        table_dims = (len(matrix),
                      max((len(row) for row in matrix), default=0))
        if dims[0] is not None and table_dims != dims:
            logger.warning(
                "merged_dimension_mismatch",
                extra={"sheet": index, "sheet_dims": dims,
                       "table_dims": table_dims},
            )
            total += fill_merged_cells(table)
            continue
        total += fill_by_ranges(table, ranges)
    return total


def filled_table_markdown(table: Any) -> str:
    """Markdown-таблица после forward-fill (тексты через ``cell_to_text``)."""
    from rag_parser.parsers.base import rows_to_markdown_table

    matrix = [
        [cell_to_text(cell) for cell in row]
        for row in (
            table.get("cells")
            if isinstance(table, dict)
            else getattr(table, "cells", [])
        )
        if isinstance(row, (list, tuple))
    ]
    filled = forward_fill_matrix(matrix)
    markdown, _, _ = rows_to_markdown_table(filled)
    return markdown


__all__ = [
    "fill_by_ranges",
    "fill_document_tables",
    "fill_merged_cells",
    "fill_workbook_tables",
    "filled_table_markdown",
    "forward_fill_matrix",
    "is_empty_value",
    "normalize_ranges",
    "sheet_merged_ranges",
]
